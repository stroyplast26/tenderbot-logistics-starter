"""Bounded, bytes-only import batches for the offline Source Lab.

This module deliberately has no file opener, URL client, credential field, or
background worker.  A caller supplies already-authorised bytes plus a typed
snapshot of the passport/access/evidence decision made outside this importer.
Every row is parsed and validated before the first Source Lab transaction.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import re
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta, timezone
from enum import Enum
from typing import Any, Callable, Mapping, Protocol, Sequence, runtime_checkable
from urllib.parse import urlsplit

from .ids import payload_hash
from .source_lab import (
    SourceLabBatchRecord,
    canonical_identity_fingerprints,
    strict_json_dumps,
    validate_source_payload,
)


SOURCE_IMPORT_PARSER_VERSION = "source-import-parser-v2"
SOURCE_IMPORT_RECORD_VERSION = "source-import-record-v2"


class SourceImportError(RuntimeError):
    """Base error whose public message never contains source cell values."""


class SourceImportValidationError(SourceImportError):
    """The offline fixture or its trust contract is invalid."""


class SourceImportConflict(SourceImportError):
    """An immutable batch/idempotency identity was reused with other bytes."""


class SourceImportSinkError(SourceImportError):
    """The local Source Lab sink failed without exposing its exception text."""


class SourceImportFormat(str, Enum):
    CSV = "CSV"
    XLSX = "XLSX"
    JSON = "JSON"
    JSONL = "JSONL"


class BomPolicy(str, Enum):
    FORBID = "FORBID"
    ALLOW = "ALLOW"
    REQUIRE = "REQUIRE"


@dataclass(frozen=True, slots=True, repr=False)
class SourceAuthorizationSnapshot:
    """Content-bound result of a separate passport/access decision.

    The importer does not create or approve this snapshot.  Production callers
    are expected to populate the references from the persistent SourcePassport,
    SourceAccess permit and evidence receipt ledgers.  The exact snapshot hash
    is included in the immutable import manifest.
    """

    snapshot_version: str
    source_id: str
    data_class: str
    acquisition_mode: str
    passport_id: str
    passport_version: str
    passport_evidence_ref: str
    access_permit_id: str
    access_policy_version: str
    access_evidence_ref: str
    evidence_receipt_id: str
    source_blob_evidence_ref: str
    content_sha256: str
    byte_count: int
    record_count: int
    captured_at_utc: str
    valid_from_utc: str
    valid_until_utc: str
    source_read_epoch: int

    def __repr__(self) -> str:
        return "SourceAuthorizationSnapshot(<redacted>)"


@dataclass(frozen=True, slots=True, repr=False)
class FieldMapping:
    source_header: str
    canonical_field: str

    def __repr__(self) -> str:
        return "FieldMapping(<redacted>)"


@dataclass(frozen=True, slots=True, repr=False)
class IdentityMapping:
    source_header: str
    namespace: str
    required: bool = False

    def __repr__(self) -> str:
        return "IdentityMapping(<redacted>)"


@dataclass(frozen=True, slots=True, repr=False)
class SourceImportPolicy:
    policy_id: str
    policy_version: str
    evidence_ref: str
    source_id: str
    acquisition_mode: str
    data_class: str
    data_contract_version: str
    allowed_formats: tuple[SourceImportFormat | str, ...]
    allowed_source_headers: tuple[str, ...]
    required_source_headers: tuple[str, ...]
    external_key_header: str
    field_mappings: tuple[FieldMapping, ...]
    identity_mappings: tuple[IdentityMapping, ...]
    authorization: SourceAuthorizationSnapshot
    text_encoding: str = "utf-8"
    bom_policy: BomPolicy | str = BomPolicy.FORBID
    csv_delimiter: str = ","
    csv_quotechar: str = '"'
    xlsx_sheet_name: str = "Sheet1"

    def __repr__(self) -> str:
        return "SourceImportPolicy(<redacted>)"

    @property
    def policy_hash(self) -> str:
        """Canonical public fingerprint of this exact mapping/trust policy."""

        return _validate_policy_shape(self).policy_hash


@dataclass(frozen=True, slots=True)
class SourceImportLimits:
    max_bytes: int = 4 * 1024 * 1024
    max_rows: int = 10_000
    max_columns: int = 128
    max_cell_bytes: int = 64 * 1024
    max_xlsx_entries: int = 512
    max_xlsx_uncompressed_bytes: int = 16 * 1024 * 1024


@dataclass(frozen=True, slots=True, repr=False)
class SourceImportResult:
    manifest_hash: str
    content_sha256: str
    mapping_policy_hash: str
    authorization_hash: str
    row_hashes: tuple[str, ...]
    source_record_ids: tuple[str, ...]
    accepted_rows: int
    created_rows: int
    replayed_rows: int
    record_results: tuple[Any, ...]

    def __repr__(self) -> str:
        return (
            "SourceImportResult(accepted_rows="
            f"{self.accepted_rows}, created_rows={self.created_rows}, "
            f"replayed_rows={self.replayed_rows}, hashes=<redacted>)"
        )

    @property
    def policy_hash(self) -> str:
        """Backward-friendly alias for the mapping-policy fingerprint."""

        return self.mapping_policy_hash


@runtime_checkable
class SourceImportSink(Protocol):
    def ingest_batch(self, **kwargs: Any) -> Sequence[Any]: ...


_SAFE_ID = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_.:/-]{0,255}$")
_CANONICAL_FIELD = re.compile(r"^[a-z][a-z0-9_]{0,63}$")
_IDENTITY_NAMESPACE = re.compile(r"^[a-z][a-z0-9_.-]{0,63}$")
_HEX64 = re.compile(r"^[0-9a-f]{64}$")
_ALLOWED_MODES = frozenset({"OFFLINE_FIXTURE"})
_ALLOWED_ENCODINGS = frozenset({"utf-8", "windows-1251"})
_UTF8_BOM = b"\xef\xbb\xbf"
_OTHER_BOMS = (b"\xff\xfe\x00\x00", b"\x00\x00\xfe\xff", b"\xff\xfe", b"\xfe\xff")
_HARD_LIMITS = SourceImportLimits()
_FORBIDDEN_XLSX_PART_PREFIXES = (
    "xl/macrosheets/",
    "xl/dialogsheets/",
    "xl/activex/",
    "xl/embeddings/",
    "xl/externallinks/",
    "xl/querytables/",
    "customui/",
)
_FORBIDDEN_XLSX_PART_NAMES = frozenset(
    {
        "xl/connections.xml",
        "xl/vbaproject.bin",
    }
)
_FORBIDDEN_OOXML_RELATION_TOKENS = (
    "macrosheet",
    "dialogsheet",
    "vbaproject",
    "activex",
    "oleobject",
    "externallink",
    "connections",
    "querytable",
    "control",
)
_FORBIDDEN_OOXML_CONTENT_TYPE_TOKENS = (
    "macrosheet",
    "dialogsheet",
    "vbaproject",
    "activex",
    "oleobject",
    "externallink",
    "connections",
    "querytable",
    "control",
    "macroenabled",
)


@dataclass(frozen=True, slots=True)
class _ValidatedPolicy:
    source_id: str
    mode: str
    data_class: str
    data_contract_version: str
    formats: frozenset[SourceImportFormat]
    allowed_headers: frozenset[str]
    required_headers: frozenset[str]
    external_header: str
    field_mappings: tuple[FieldMapping, ...]
    identity_mappings: tuple[IdentityMapping, ...]
    encoding: str
    bom: BomPolicy
    delimiter: str
    quotechar: str
    sheet_name: str
    policy_hash: str
    authorization_hash: str
    captured_at_utc: str
    evidence_ref: str
    policy_body: Mapping[str, Any]
    authorization_body: Mapping[str, Any]


@dataclass(frozen=True, slots=True)
class _PreparedRow:
    external_key: str
    payload: Mapping[str, Any]
    canonical_keys: tuple[tuple[str, str], ...]
    row_hash: str


def _required(value: object, message: str, *, limit: int = 512) -> str:
    if not isinstance(value, str):
        raise SourceImportValidationError(message)
    text = value.strip()
    try:
        encoded = text.encode("utf-8", "strict")
    except UnicodeError:
        raise SourceImportValidationError(message) from None
    if not text or len(text) > limit or len(encoded) > limit * 4:
        raise SourceImportValidationError(message)
    return text


def _safe_id(value: object, message: str) -> str:
    text = _required(value, message, limit=256)
    if not _SAFE_ID.fullmatch(text):
        raise SourceImportValidationError(message)
    return text


def _evidence_ref(value: object, message: str) -> str:
    text = _required(value, message, limit=2048)
    try:
        parsed = urlsplit(text)
    except ValueError:
        raise SourceImportValidationError(message) from None
    if (
        parsed.scheme not in {"evidence", "radar-evidence", "source-evidence"}
        or not (parsed.netloc or parsed.path)
        or parsed.username is not None
        or parsed.password is not None
        or parsed.query
        or parsed.fragment
        or any(character.isspace() for character in text)
    ):
        raise SourceImportValidationError(message)
    return text


def _timestamp(value: object, message: str) -> tuple[str, datetime]:
    raw = _required(value, message, limit=64)
    try:
        parsed = datetime.fromisoformat(raw.replace("Z", "+00:00"))
    except (TypeError, ValueError):
        raise SourceImportValidationError(message) from None
    if parsed.tzinfo is None or parsed.utcoffset() is None:
        raise SourceImportValidationError(message)
    parsed = parsed.astimezone(timezone.utc)
    rendered = parsed.isoformat(timespec="microseconds").replace("+00:00", "Z")
    return rendered.replace(".000000Z", "Z"), parsed


def _strict_json(value: object, message: str) -> str:
    try:
        rendered = strict_json_dumps(value)
    except Exception:
        raise SourceImportValidationError(message) from None
    return rendered


def _sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def _bounded_limits(limits: SourceImportLimits) -> SourceImportLimits:
    if not isinstance(limits, SourceImportLimits):
        raise SourceImportValidationError("source import limits are invalid")
    values = (
        limits.max_bytes,
        limits.max_rows,
        limits.max_columns,
        limits.max_cell_bytes,
        limits.max_xlsx_entries,
        limits.max_xlsx_uncompressed_bytes,
    )
    ceilings = (
        _HARD_LIMITS.max_bytes,
        _HARD_LIMITS.max_rows,
        _HARD_LIMITS.max_columns,
        _HARD_LIMITS.max_cell_bytes,
        _HARD_LIMITS.max_xlsx_entries,
        _HARD_LIMITS.max_xlsx_uncompressed_bytes,
    )
    if any(isinstance(item, bool) for item in values) or any(
        not isinstance(item, int) or item < 1 or item > ceiling
        for item, ceiling in zip(values, ceilings)
    ):
        raise SourceImportValidationError("source import limits are invalid")
    return limits


def _header_tuple(values: object, message: str) -> tuple[str, ...]:
    if not isinstance(values, tuple) or not values:
        raise SourceImportValidationError(message)
    result: list[str] = []
    for value in values:
        header = _required(value, message, limit=128)
        if header != value or "\x00" in header:
            raise SourceImportValidationError(message)
        result.append(header)
    if len(set(result)) != len(result):
        raise SourceImportValidationError(message)
    return tuple(result)


def _authorization_body(auth: SourceAuthorizationSnapshot) -> dict[str, Any]:
    return {
        "snapshot_version": auth.snapshot_version,
        "source_id": auth.source_id,
        "data_class": auth.data_class,
        "acquisition_mode": auth.acquisition_mode,
        "passport_id": auth.passport_id,
        "passport_version": auth.passport_version,
        "passport_evidence_ref": auth.passport_evidence_ref,
        "access_permit_id": auth.access_permit_id,
        "access_policy_version": auth.access_policy_version,
        "access_evidence_ref": auth.access_evidence_ref,
        "evidence_receipt_id": auth.evidence_receipt_id,
        "source_blob_evidence_ref": auth.source_blob_evidence_ref,
        "content_sha256": auth.content_sha256,
        "byte_count": auth.byte_count,
        "record_count": auth.record_count,
        "captured_at_utc": auth.captured_at_utc,
        "valid_from_utc": auth.valid_from_utc,
        "valid_until_utc": auth.valid_until_utc,
        "source_read_epoch": auth.source_read_epoch,
    }


def _validate_policy(
    policy: SourceImportPolicy,
    *,
    content_sha256: str,
    byte_count: int,
    record_count: int,
    now: datetime,
    current_source_read_epoch: int,
) -> _ValidatedPolicy:
    if not isinstance(policy, SourceImportPolicy):
        raise SourceImportValidationError("source import policy is invalid")
    policy_id = _safe_id(policy.policy_id, "source import policy identity is invalid")
    policy_version = _safe_id(
        policy.policy_version, "source import policy version is invalid"
    )
    policy_evidence = _evidence_ref(
        policy.evidence_ref, "source import policy evidence is required"
    )
    source_id = _safe_id(policy.source_id, "source import policy source is invalid")
    data_class = _safe_id(policy.data_class, "source import data class is invalid").upper()
    data_contract_version = _safe_id(
        policy.data_contract_version, "source import data contract version is invalid"
    )
    mode = _safe_id(policy.acquisition_mode, "source import mode is invalid").upper()
    if mode not in _ALLOWED_MODES:
        raise SourceImportValidationError("source import mode is not supported by the persistent ledger")
    if not isinstance(policy.allowed_formats, tuple) or not policy.allowed_formats:
        raise SourceImportValidationError("source import formats are invalid")
    try:
        formats = frozenset(
            item if isinstance(item, SourceImportFormat) else SourceImportFormat(str(item).upper())
            for item in policy.allowed_formats
        )
    except ValueError:
        raise SourceImportValidationError("source import formats are invalid") from None
    if len(formats) != len(policy.allowed_formats):
        raise SourceImportValidationError("source import formats are invalid")

    allowed = _header_tuple(
        policy.allowed_source_headers, "source import allowed headers are invalid"
    )
    required = _header_tuple(
        policy.required_source_headers, "source import required headers are invalid"
    )
    allowed_set, required_set = frozenset(allowed), frozenset(required)
    if len(allowed) > _HARD_LIMITS.max_columns or not required_set.issubset(allowed_set):
        raise SourceImportValidationError("source import header contract is invalid")
    external_header = _required(
        policy.external_key_header, "source import external key mapping is invalid", limit=128
    )
    if external_header not in required_set:
        raise SourceImportValidationError("source import external key mapping is invalid")

    if not isinstance(policy.field_mappings, tuple) or not policy.field_mappings:
        raise SourceImportValidationError("source import field mappings are invalid")
    mapped_sources: set[str] = set()
    mapped_targets: set[str] = set()
    for mapping in policy.field_mappings:
        if not isinstance(mapping, FieldMapping):
            raise SourceImportValidationError("source import field mappings are invalid")
        source_header = _required(
            mapping.source_header, "source import field mappings are invalid", limit=128
        )
        target = _required(
            mapping.canonical_field, "source import field mappings are invalid", limit=64
        )
        if (
            source_header not in allowed_set
            or not _CANONICAL_FIELD.fullmatch(target)
            or source_header in mapped_sources
            or target in mapped_targets
        ):
            raise SourceImportValidationError("source import field mappings are invalid")
        mapped_sources.add(source_header)
        mapped_targets.add(target)

    if not isinstance(policy.identity_mappings, tuple):
        raise SourceImportValidationError("source import identity mappings are invalid")
    identity_pairs: set[tuple[str, str]] = set()
    for mapping in policy.identity_mappings:
        if not isinstance(mapping, IdentityMapping) or not isinstance(mapping.required, bool):
            raise SourceImportValidationError("source import identity mappings are invalid")
        source_header = _required(
            mapping.source_header, "source import identity mappings are invalid", limit=128
        )
        namespace = _required(
            mapping.namespace, "source import identity mappings are invalid", limit=64
        ).lower()
        pair = (source_header, namespace)
        if (
            source_header not in allowed_set
            or source_header not in mapped_sources
            or not _IDENTITY_NAMESPACE.fullmatch(namespace)
            or pair in identity_pairs
        ):
            raise SourceImportValidationError("source import identity mappings are invalid")
        identity_pairs.add(pair)

    encoding = str(policy.text_encoding or "").strip().lower()
    if encoding not in _ALLOWED_ENCODINGS:
        raise SourceImportValidationError("source import text encoding is invalid")
    try:
        bom = policy.bom_policy if isinstance(policy.bom_policy, BomPolicy) else BomPolicy(
            str(policy.bom_policy).upper()
        )
    except ValueError:
        raise SourceImportValidationError("source import BOM policy is invalid") from None
    if encoding != "utf-8" and bom is not BomPolicy.FORBID:
        raise SourceImportValidationError("source import BOM policy is invalid")
    delimiter = str(policy.csv_delimiter)
    quotechar = str(policy.csv_quotechar)
    if delimiter not in {",", ";", "|", "\t"} or quotechar not in {'"', "'"}:
        raise SourceImportValidationError("source import CSV dialect is invalid")
    if delimiter == quotechar:
        raise SourceImportValidationError("source import CSV dialect is invalid")
    sheet_name = _required(
        policy.xlsx_sheet_name, "source import workbook sheet is invalid", limit=64
    )

    auth = policy.authorization
    if not isinstance(auth, SourceAuthorizationSnapshot):
        raise SourceImportValidationError("source import authorization is required")
    _safe_id(auth.snapshot_version, "source import authorization version is invalid")
    auth_source = _safe_id(auth.source_id, "source import authorization source is invalid")
    auth_class = _safe_id(
        auth.data_class, "source import authorization data class is invalid"
    ).upper()
    auth_mode = _safe_id(auth.acquisition_mode, "source import authorization mode is invalid").upper()
    if auth_mode not in _ALLOWED_MODES or (auth_source, auth_class, auth_mode) != (
        source_id,
        data_class,
        mode,
    ):
        raise SourceImportValidationError("source import authorization scope is invalid")
    for value, message in (
        (auth.passport_id, "source passport identity is required"),
        (auth.passport_version, "source passport version is required"),
        (auth.access_permit_id, "source access permit identity is required"),
        (auth.access_policy_version, "source access policy version is required"),
        (auth.evidence_receipt_id, "source evidence receipt is required"),
    ):
        _safe_id(value, message)
    if auth.snapshot_version != "source-authorization-v1" or auth.access_policy_version != "source-access-permit-v1":
        raise SourceImportValidationError("source import authorization version is invalid")
    for value, message in (
        (auth.passport_evidence_ref, "source passport evidence is required"),
        (auth.access_evidence_ref, "source access evidence is required"),
        (auth.source_blob_evidence_ref, "source blob evidence is required"),
    ):
        _evidence_ref(value, message)
    declared_hash = str(auth.content_sha256 or "").strip().lower()
    if not _HEX64.fullmatch(declared_hash):
        raise SourceImportValidationError("source authorization content hash is invalid")
    for number in (auth.byte_count, auth.record_count):
        if isinstance(number, bool) or not isinstance(number, int) or number < 1:
            raise SourceImportValidationError("source authorization usage is invalid")
    if (
        isinstance(auth.source_read_epoch, bool)
        or not isinstance(auth.source_read_epoch, int)
        or not 0 <= auth.source_read_epoch < 10**32
    ):
        raise SourceImportValidationError("source authorization read epoch is invalid")
    if auth.source_read_epoch != current_source_read_epoch:
        raise SourceImportValidationError("source authorization belongs to an obsolete read epoch")
    if (
        declared_hash != content_sha256
        or auth.byte_count != byte_count
        or auth.record_count != record_count
    ):
        raise SourceImportValidationError("source authorization does not match the fixture")
    captured_text, captured = _timestamp(
        auth.captured_at_utc, "source authorization timestamp is invalid"
    )
    _, valid_from = _timestamp(auth.valid_from_utc, "source authorization validity is invalid")
    _, valid_until = _timestamp(auth.valid_until_utc, "source authorization validity is invalid")
    if valid_until < valid_from or not (valid_from <= captured <= valid_until):
        raise SourceImportValidationError("source authorization validity is invalid")
    if not (valid_from <= now <= valid_until) or captured > now + timedelta(minutes=5):
        raise SourceImportValidationError("source import authorization is not currently valid")

    auth_body = _authorization_body(auth)
    _strict_json(auth_body, "source import authorization is invalid")
    authorization_hash = payload_hash(auth_body)
    policy_body = {
        "policy_schema_version": "source-import-policy-v1",
        "parser_version": SOURCE_IMPORT_PARSER_VERSION,
        "policy_id": policy_id,
        "policy_version": policy_version,
        "evidence_ref": policy_evidence,
        "source_id": source_id,
        "acquisition_mode": mode,
        "data_class": data_class,
        "data_contract_version": data_contract_version,
        "allowed_formats": sorted(item.value for item in formats),
        "allowed_source_headers": list(allowed),
        "required_source_headers": list(required),
        "external_key_header": external_header,
        "field_mappings": [
            [item.source_header, item.canonical_field] for item in policy.field_mappings
        ],
        "identity_mappings": [
            [item.source_header, item.namespace.lower(), item.required]
            for item in policy.identity_mappings
        ],
        "text_encoding": encoding,
        "bom_policy": bom.value,
        "csv_delimiter": delimiter,
        "csv_quotechar": quotechar,
        "xlsx_sheet_name": sheet_name,
    }
    _strict_json(policy_body, "source import policy is invalid")
    return _ValidatedPolicy(
        source_id,
        mode,
        data_class,
        data_contract_version,
        formats,
        allowed_set,
        required_set,
        external_header,
        policy.field_mappings,
        policy.identity_mappings,
        encoding,
        bom,
        delimiter,
        quotechar,
        sheet_name,
        payload_hash(policy_body),
        authorization_hash,
        captured_text,
        auth.source_blob_evidence_ref,
        policy_body,
        auth_body,
    )


def _decode_text(
    data: bytes,
    policy: _ValidatedPolicy,
    *,
    force_utf8: bool = False,
) -> str:
    if any(data.startswith(marker) for marker in _OTHER_BOMS):
        raise SourceImportValidationError("source import text BOM is invalid")
    has_bom = data.startswith(_UTF8_BOM)
    if policy.bom is BomPolicy.FORBID and has_bom:
        raise SourceImportValidationError("source import text BOM is invalid")
    if policy.bom is BomPolicy.REQUIRE and not has_bom:
        raise SourceImportValidationError("source import text BOM is invalid")
    raw = data[len(_UTF8_BOM) :] if has_bom else data
    codec = "utf-8" if force_utf8 or policy.encoding == "utf-8" else "cp1251"
    try:
        text = raw.decode(codec, "strict")
        text.encode("utf-8", "strict")
    except UnicodeError:
        raise SourceImportValidationError("source import text encoding is invalid") from None
    if "\x00" in text:
        raise SourceImportValidationError("source import text contains invalid characters")
    return text


def _check_headers(headers: Sequence[object], policy: _ValidatedPolicy, limits: SourceImportLimits) -> tuple[str, ...]:
    if not headers or len(headers) > limits.max_columns:
        raise SourceImportValidationError("source import header row is invalid")
    normalized: list[str] = []
    for value in headers:
        if not isinstance(value, str) or not value or value != value.strip():
            raise SourceImportValidationError("source import header row is invalid")
        if len(value.encode("utf-8", "strict")) > limits.max_cell_bytes:
            raise SourceImportValidationError("source import header row is invalid")
        normalized.append(value)
    if len(set(normalized)) != len(normalized):
        raise SourceImportValidationError("source import header row is invalid")
    actual = frozenset(normalized)
    if not actual.issubset(policy.allowed_headers) or not policy.required_headers.issubset(actual):
        raise SourceImportValidationError("source import header contract does not match")
    return tuple(normalized)


def _cell_json_size(value: object, limits: SourceImportLimits) -> None:
    rendered = _strict_json(value, "source import cell type is invalid")
    if len(rendered.encode("utf-8", "strict")) > limits.max_cell_bytes:
        raise SourceImportValidationError("source import cell exceeds the limit")


def _validate_objects(
    rows: Sequence[Mapping[str, Any]], policy: _ValidatedPolicy, limits: SourceImportLimits
) -> list[dict[str, Any]]:
    if not rows:
        raise SourceImportValidationError("source import contains no records")
    if len(rows) > limits.max_rows:
        raise SourceImportValidationError("source import row limit exceeded")
    validated: list[dict[str, Any]] = []
    for row in rows:
        if not isinstance(row, Mapping) or not row or len(row) > limits.max_columns:
            raise SourceImportValidationError("source import record shape is invalid")
        if any(not isinstance(key, str) for key in row):
            raise SourceImportValidationError("source import record keys are invalid")
        keys = frozenset(row)
        if not keys.issubset(policy.allowed_headers) or not policy.required_headers.issubset(keys):
            raise SourceImportValidationError("source import record header contract does not match")
        normalized = dict(row)
        for value in normalized.values():
            _cell_json_size(value, limits)
        validated.append(normalized)
    return validated


def _parse_csv(data: bytes, policy: _ValidatedPolicy, limits: SourceImportLimits) -> list[dict[str, Any]]:
    text = _decode_text(data, policy)
    try:
        rows = list(
            csv.reader(
                io.StringIO(text, newline=""),
                delimiter=policy.delimiter,
                quotechar=policy.quotechar,
                doublequote=True,
                strict=True,
            )
        )
    except csv.Error:
        raise SourceImportValidationError("source import CSV syntax is invalid") from None
    if not rows:
        raise SourceImportValidationError("source import contains no header")
    headers = _check_headers(rows[0], policy, limits)
    objects: list[dict[str, Any]] = []
    for values in rows[1:]:
        if len(values) != len(headers) or not any(value != "" for value in values):
            raise SourceImportValidationError("source import CSV row shape is invalid")
        objects.append(dict(zip(headers, values)))
        if len(objects) > limits.max_rows:
            raise SourceImportValidationError("source import row limit exceeded")
    return _validate_objects(objects, policy, limits)


def _json_pairs(pairs: list[tuple[str, Any]]) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for key, value in pairs:
        if key in result:
            raise SourceImportValidationError("source import JSON object keys are duplicated")
        result[key] = value
    return result


def _json_constant(_value: str) -> None:
    raise SourceImportValidationError("source import JSON number is invalid")


def _parse_one_json(text: str) -> Any:
    try:
        return json.loads(
            text,
            object_pairs_hook=_json_pairs,
            parse_constant=_json_constant,
        )
    except SourceImportValidationError:
        raise
    except (json.JSONDecodeError, TypeError, ValueError, RecursionError):
        raise SourceImportValidationError("source import JSON syntax is invalid") from None


def _parse_json(data: bytes, policy: _ValidatedPolicy, limits: SourceImportLimits, *, lines: bool) -> list[dict[str, Any]]:
    # RFC 8259 interchange is UTF-8.  The legacy CSV codec option must never
    # weaken JSON/JSONL Unicode validation.
    text = _decode_text(data, policy, force_utf8=True)
    if lines:
        raw_lines = text.splitlines()
        if not raw_lines or any(not line.strip() for line in raw_lines):
            raise SourceImportValidationError("source import JSONL framing is invalid")
        if len(raw_lines) > limits.max_rows:
            raise SourceImportValidationError("source import row limit exceeded")
        values = [_parse_one_json(line) for line in raw_lines]
    else:
        value = _parse_one_json(text)
        if not isinstance(value, list):
            raise SourceImportValidationError("source import JSON root must be an array")
        values = value
    return _validate_objects(values, policy, limits)


def _xlsx_scalar(value: object) -> object:
    if value is None or isinstance(value, (str, bool, int)):
        return value
    if isinstance(value, float):
        if not math.isfinite(value):
            raise SourceImportValidationError("source import workbook cell is invalid")
        return value
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    raise SourceImportValidationError("source import workbook cell type is invalid")


def _xlsx_part_is_forbidden(name: str) -> bool:
    normalized = str(name or "").replace("\\", "/").lstrip("/").lower()
    return normalized in _FORBIDDEN_XLSX_PART_NAMES or normalized.startswith(
        _FORBIDDEN_XLSX_PART_PREFIXES
    )


def _inspect_ooxml_control_xml(body: bytes, *, relationships: bool) -> None:
    """Reject active/external OOXML declarations after XML entity decoding."""

    lowered = body.lower()
    if b"<!doctype" in lowered or b"<!entity" in lowered:
        raise SourceImportValidationError("source import workbook active content is forbidden")
    try:
        root = ET.fromstring(body)
    except (ET.ParseError, ValueError, RecursionError):
        raise SourceImportValidationError("source import workbook archive is invalid") from None
    if relationships:
        for element in root.iter():
            relation_type = str(element.attrib.get("Type", "")).strip().lower()
            target_mode = str(element.attrib.get("TargetMode", "")).strip().lower()
            target = str(element.attrib.get("Target", "")).strip().replace("\\", "/")
            target_lower = target.lower()
            if (
                target_mode == "external"
                or any(token in relation_type for token in _FORBIDDEN_OOXML_RELATION_TOKENS)
                or _xlsx_part_is_forbidden(target_lower)
                or "://" in target_lower
            ):
                raise SourceImportValidationError(
                    "source import workbook active or external relationships are forbidden"
                )
    else:
        for element in root.iter():
            content_type = str(element.attrib.get("ContentType", "")).strip().lower()
            part_name = str(element.attrib.get("PartName", "")).strip()
            if (
                any(
                    token in content_type
                    for token in _FORBIDDEN_OOXML_CONTENT_TYPE_TOKENS
                )
                or _xlsx_part_is_forbidden(part_name)
            ):
                raise SourceImportValidationError(
                    "source import workbook active content is forbidden"
                )


def _inspect_xlsx_archive(data: bytes, limits: SourceImportLimits) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            names = [item.filename.replace("\\", "/").lower() for item in entries]
            if (
                len(entries) > limits.max_xlsx_entries
                or len(set(names)) != len(names)
                or sum(item.file_size for item in entries) > limits.max_xlsx_uncompressed_bytes
                or any(item.flag_bits & 0x1 for item in entries)
            ):
                raise SourceImportValidationError("source import workbook archive is unsafe")
            if any(
                _xlsx_part_is_forbidden(name)
                for name in names
            ):
                raise SourceImportValidationError("source import workbook active content is forbidden")
            for item, name in zip(entries, names):
                if item.file_size > limits.max_xlsx_uncompressed_bytes:
                    raise SourceImportValidationError("source import workbook archive is unsafe")
                if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                    body = archive.read(item)
                    lowered = body.lower()
                    if (
                        re.search(br"<(?:[a-z0-9_]+:)?f(?:\s|>)", lowered)
                        or b'hidden="1"' in lowered
                        or b'hidden="true"' in lowered
                        or b"<mergecells" in lowered
                    ):
                        raise SourceImportValidationError("source import workbook contains ambiguous cells")
                if name == "[content_types].xml":
                    _inspect_ooxml_control_xml(archive.read(item), relationships=False)
                if name.endswith(".rels"):
                    _inspect_ooxml_control_xml(archive.read(item), relationships=True)
            if archive.testzip() is not None:
                raise SourceImportValidationError("source import workbook archive is corrupt")
    except SourceImportValidationError:
        raise
    except (zipfile.BadZipFile, OSError, RuntimeError):
        raise SourceImportValidationError("source import workbook archive is invalid") from None


def _parse_xlsx(data: bytes, policy: _ValidatedPolicy, limits: SourceImportLimits) -> list[dict[str, Any]]:
    _inspect_xlsx_archive(data, limits)
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError:
        raise SourceImportValidationError("source import workbook support is unavailable") from None
    try:
        workbook = load_workbook(
            io.BytesIO(data), read_only=True, data_only=True, keep_links=True
        )
    except (InvalidFileException, OSError, ValueError, KeyError, zipfile.BadZipFile):
        raise SourceImportValidationError("source import workbook cannot be read") from None
    try:
        if (
            len(workbook.worksheets) != 1
            or workbook.worksheets[0].title != policy.sheet_name
            or workbook.worksheets[0].sheet_state != "visible"
            or getattr(workbook, "_external_links", ())
        ):
            raise SourceImportValidationError("source import workbook sheet contract does not match")
        sheet = workbook.worksheets[0]
        if sheet.max_row > limits.max_rows + 1 or sheet.max_column > limits.max_columns:
            raise SourceImportValidationError("source import workbook dimensions exceed the limit")
        raw_rows: list[list[object]] = []
        for row in sheet.iter_rows(
            min_row=1,
            max_row=limits.max_rows + 2,
            max_col=limits.max_columns + 1,
        ):
            values: list[object] = []
            for cell in row:
                if getattr(cell, "data_type", "") in {"f", "e"}:
                    raise SourceImportValidationError("source import workbook cell is invalid")
                values.append(_xlsx_scalar(cell.value))
            while values and values[-1] is None:
                values.pop()
            if values:
                raw_rows.append(values)
        if not raw_rows:
            raise SourceImportValidationError("source import contains no header")
        headers = _check_headers(raw_rows[0], policy, limits)
        objects: list[dict[str, Any]] = []
        for values in raw_rows[1:]:
            if len(values) > len(headers):
                raise SourceImportValidationError("source import workbook row shape is invalid")
            values.extend([None] * (len(headers) - len(values)))
            if not any(value is not None and value != "" for value in values):
                raise SourceImportValidationError("source import workbook row shape is invalid")
            objects.append(dict(zip(headers, values)))
        return _validate_objects(objects, policy, limits)
    finally:
        workbook.close()


def _external_key(value: object) -> str:
    if isinstance(value, bool) or not isinstance(value, (str, int)):
        raise SourceImportValidationError("source import external key is invalid")
    text = str(value).strip()
    if not text or len(text) > 1024:
        raise SourceImportValidationError("source import external key is invalid")
    return text


def _identity_value(value: object, *, required: bool) -> str:
    if value is None or value == "":
        if required:
            raise SourceImportValidationError("source import canonical identity is missing")
        return ""
    if isinstance(value, bool) or not isinstance(value, (str, int)):
        raise SourceImportValidationError("source import canonical identity is invalid")
    text = str(value).strip()
    if not text and required:
        raise SourceImportValidationError("source import canonical identity is missing")
    if len(text) > 512:
        raise SourceImportValidationError("source import canonical identity is invalid")
    return text


class SourceBatchImporter:
    """Prevalidate an authorised fixture and append its rows to Source Lab."""

    def __init__(
        self,
        sink: SourceImportSink,
        *,
        policy: SourceImportPolicy,
        limits: SourceImportLimits = SourceImportLimits(),
        clock: Callable[[], datetime] | None = None,
        current_source_read_epoch: int | None = None,
    ) -> None:
        if not isinstance(sink, SourceImportSink):
            raise SourceImportValidationError("source import sink is invalid")
        self._sink = sink
        self._policy = policy
        self._limits = _bounded_limits(limits)
        self._clock = clock or (lambda: datetime.now(timezone.utc))
        if current_source_read_epoch is not None and (
            isinstance(current_source_read_epoch, bool)
            or not isinstance(current_source_read_epoch, int)
            or not 0 <= current_source_read_epoch < 10**32
        ):
            raise SourceImportValidationError("current source read epoch is invalid")
        self._source_read_epoch = current_source_read_epoch

    def _current_source_read_epoch(self) -> int:
        if self._source_read_epoch is not None:
            return self._source_read_epoch
        store = getattr(self._sink, "store", None)
        connect = getattr(store, "connect", None)
        if not callable(connect):
            raise SourceImportValidationError("current source read epoch is required")
        con = None
        try:
            con = connect()
            row = con.execute(
                "SELECT value FROM schema_meta WHERE key='source_read_epoch'"
            ).fetchone()
            raw = str(row[0] if row else "")
            if not re.fullmatch(r"[0-9]{32}", raw):
                raise SourceImportValidationError("current source read epoch is invalid")
            return int(raw)
        except SourceImportValidationError:
            raise
        except Exception:
            raise SourceImportValidationError("current source read epoch cannot be verified") from None
        finally:
            if con is not None:
                con.close()

    def import_bytes(
        self,
        data: bytes,
        *,
        source_format: SourceImportFormat | str,
        run_key: str,
        batch_key: str,
    ) -> SourceImportResult:
        if type(data) is not bytes:
            raise SourceImportValidationError("source import requires an in-memory bytes value")
        if not data or len(data) > self._limits.max_bytes:
            raise SourceImportValidationError("source import byte limit exceeded")
        run = _safe_id(run_key, "source import run key is invalid")
        batch = _safe_id(batch_key, "source import batch key is invalid")
        try:
            fmt = source_format if isinstance(source_format, SourceImportFormat) else SourceImportFormat(
                str(source_format).upper()
            )
        except ValueError:
            raise SourceImportValidationError("source import format is invalid") from None
        now = self._clock()
        if not isinstance(now, datetime) or now.tzinfo is None or now.utcoffset() is None:
            raise SourceImportValidationError("source import clock is invalid")
        now = now.astimezone(timezone.utc)

        # Parsing precedes authorization matching because the persistent access
        # receipt is expected to bind the exact record count as well as bytes.
        # No sink call occurs anywhere in this phase.
        content_sha256 = _sha256_bytes(data)
        provisional = _validate_policy_shape(self._policy)
        if fmt not in provisional.formats:
            raise SourceImportValidationError("source import format is not allowed")
        try:
            if fmt is SourceImportFormat.CSV:
                rows = _parse_csv(data, provisional, self._limits)
            elif fmt is SourceImportFormat.XLSX:
                rows = _parse_xlsx(data, provisional, self._limits)
            elif fmt is SourceImportFormat.JSON:
                rows = _parse_json(data, provisional, self._limits, lines=False)
            else:
                rows = _parse_json(data, provisional, self._limits, lines=True)
        except SourceImportError:
            raise
        except Exception:
            raise SourceImportValidationError("source import fixture cannot be parsed safely") from None
        policy = _validate_policy(
            self._policy,
            content_sha256=content_sha256,
            byte_count=len(data),
            record_count=len(rows),
            now=now,
            current_source_read_epoch=self._current_source_read_epoch(),
        )

        prepared: list[_PreparedRow] = []
        external_keys_seen: set[str] = set()
        for index, row in enumerate(rows, start=1):
            external = _external_key(row[policy.external_header])
            if external in external_keys_seen:
                raise SourceImportValidationError(
                    "source import contains a duplicate external key"
                )
            external_keys_seen.add(external)
            mapped = {
                mapping.canonical_field: row[mapping.source_header]
                for mapping in policy.field_mappings
                if mapping.source_header in row
            }
            canonical_keys: list[tuple[str, str]] = []
            for mapping in policy.identity_mappings:
                identity = _identity_value(
                    row.get(mapping.source_header), required=mapping.required
                )
                if identity:
                    canonical_keys.append((mapping.namespace.lower(), identity))
            try:
                canonical_identity_fingerprints(canonical_keys)
            except Exception:
                raise SourceImportValidationError(
                    "source import canonical identity is invalid"
                ) from None
            row_body = {
                "source_import_row_version": 2,
                "row_number": index,
                "external_key_hash": payload_hash(
                    {
                        "source_lab_external_key_version": 1,
                        "source_id": policy.source_id,
                        "external_key": external,
                    }
                ),
                "mapped_record": mapped,
            }
            row_hash = hashlib.sha256(
                _strict_json(row_body, "source import row is not canonical JSON").encode(
                    "utf-8", "strict"
                )
            ).hexdigest()
            prepared.append(_PreparedRow(external, mapped, tuple(canonical_keys), row_hash))

        row_hashes = tuple(item.row_hash for item in prepared)
        manifest_body = {
            "source_import_manifest_version": 2,
            "parser_version": SOURCE_IMPORT_PARSER_VERSION,
            "source_id": policy.source_id,
            "acquisition_mode": policy.mode,
            "data_class": policy.data_class,
            "data_contract_version": policy.data_contract_version,
            "format": fmt.value,
            "run_key": run,
            "batch_key": batch,
            "content_sha256": content_sha256,
            "byte_count": len(data),
            "record_count": len(prepared),
            "ordered_row_hashes": list(row_hashes),
            "policy_hash": policy.policy_hash,
            "authorization_hash": policy.authorization_hash,
            "passport_id": self._policy.authorization.passport_id,
            "passport_version": self._policy.authorization.passport_version,
            "access_permit_id": self._policy.authorization.access_permit_id,
            "evidence_receipt_id": self._policy.authorization.evidence_receipt_id,
            "source_read_epoch": self._policy.authorization.source_read_epoch,
            "observed_at_utc": policy.captured_at_utc,
            "source_blob_evidence_ref": policy.evidence_ref,
        }
        manifest_hash = payload_hash(manifest_body)
        ordered_row_hashes_hash = payload_hash(row_hashes)

        commands: list[SourceLabBatchRecord] = []
        for index, item in enumerate(prepared, start=1):
            envelope = {
                "schema_version": SOURCE_IMPORT_RECORD_VERSION,
                "parser_version": SOURCE_IMPORT_PARSER_VERSION,
                "data_contract_version": policy.data_contract_version,
                "manifest_hash": manifest_hash,
                "content_sha256": content_sha256,
                "row_hash": item.row_hash,
                "row_number": index,
                "record_count": len(prepared),
                "ordered_row_hashes_hash": ordered_row_hashes_hash,
                "mapping_policy_hash": policy.policy_hash,
                "authorization_hash": policy.authorization_hash,
                "passport_id": self._policy.authorization.passport_id,
                "access_permit_id": self._policy.authorization.access_permit_id,
                "evidence_receipt_id": self._policy.authorization.evidence_receipt_id,
                "source_read_epoch": self._policy.authorization.source_read_epoch,
                "observed_at_utc": policy.captured_at_utc,
                "source_blob_evidence_ref": policy.evidence_ref,
                "record": item.payload,
            }
            if index == 1:
                envelope["batch_anchor"] = {
                    "import_manifest": manifest_body,
                    "mapping_policy": policy.policy_body,
                    "authorization_snapshot": policy.authorization_body,
                }
            idem = "source-import:" + payload_hash(
                {
                    "source_id": policy.source_id,
                    "run_key": run,
                    "batch_key": batch,
                    "manifest_hash": manifest_hash,
                    "row_number": index,
                    "row_hash": item.row_hash,
                }
            )
            try:
                validate_source_payload(envelope)
            except Exception:
                raise SourceImportValidationError(
                    "source import record exceeds the Source Lab contract"
                ) from None
            commands.append(
                SourceLabBatchRecord(
                    external_key=item.external_key,
                    payload=envelope,
                    observed_at_utc=policy.captured_at_utc,
                    evidence_ref=policy.evidence_ref,
                    idempotency_key=idem,
                    canonical_keys=tuple(item.canonical_keys),
                )
            )

        # Everything above is validation/canonicalisation.  The Source Lab
        # batch boundary commits every row or rolls the whole batch back.
        try:
            results = list(
                self._sink.ingest_batch(
                    source_id=policy.source_id,
                    acquisition_mode=policy.mode,
                    run_key=run,
                    batch_key=batch,
                    manifest_hash=manifest_hash,
                    records=tuple(commands),
                )
            )
        except SourceImportConflict:
            raise
        except Exception as exc:
            if exc.__class__.__name__.lower().endswith("conflict"):
                raise SourceImportConflict(
                    "source import batch identity conflicts with stored facts"
                ) from None
            raise SourceImportSinkError("source import sink failed") from None
        if len(results) != len(commands):
            raise SourceImportSinkError("source import sink returned an incomplete batch result")
        created = sum(bool(getattr(item, "created", False)) for item in results)
        return SourceImportResult(
            manifest_hash,
            content_sha256,
            policy.policy_hash,
            policy.authorization_hash,
            row_hashes,
            tuple(str(getattr(item, "source_record_id", "")) for item in results),
            len(results),
            created,
            len(results) - created,
            tuple(results),
        )


def _validate_policy_shape(policy: SourceImportPolicy) -> _ValidatedPolicy:
    """Validate parsing-related policy fields without trusting fixture claims.

    Authorization content/count/time checks happen after parsing, but this
    helper intentionally invokes the same structural validator with declared
    snapshot facts and a time inside its declared validity window.
    """

    if not isinstance(policy, SourceImportPolicy) or not isinstance(
        policy.authorization, SourceAuthorizationSnapshot
    ):
        raise SourceImportValidationError("source import policy is invalid")
    auth = policy.authorization
    _, valid_from = _timestamp(auth.valid_from_utc, "source authorization validity is invalid")
    _, valid_until = _timestamp(auth.valid_until_utc, "source authorization validity is invalid")
    if valid_until < valid_from:
        raise SourceImportValidationError("source authorization validity is invalid")
    _, captured = _timestamp(auth.captured_at_utc, "source authorization timestamp is invalid")
    return _validate_policy(
        policy,
        content_sha256=str(auth.content_sha256 or "").strip().lower(),
        byte_count=auth.byte_count,
        record_count=auth.record_count,
        now=captured,
        current_source_read_epoch=auth.source_read_epoch,
    )


def import_source_bytes(
    sink: SourceImportSink,
    data: bytes,
    *,
    policy: SourceImportPolicy,
    source_format: SourceImportFormat | str,
    run_key: str,
    batch_key: str,
    limits: SourceImportLimits = SourceImportLimits(),
    clock: Callable[[], datetime] | None = None,
    current_source_read_epoch: int | None = None,
) -> SourceImportResult:
    """Convenience wrapper; still accepts bytes only and performs no read."""

    return SourceBatchImporter(
        sink,
        policy=policy,
        limits=limits,
        clock=clock,
        current_source_read_epoch=current_source_read_epoch,
    ).import_bytes(data, source_format=source_format, run_key=run_key, batch_key=batch_key)
