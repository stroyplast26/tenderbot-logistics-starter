from __future__ import annotations

import argparse
import re
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = ROOT / "reports" / "market_demand_os_deep_research_2026-08-25" / "report-source.md"
DEFAULT_OUTPUT = ROOT / "reports" / "market_demand_os_deep_research_2026-08-25" / "АлюмКомплект_Market_Demand_OS_v7_2026-08-25.docx"
DEFAULT_ASSET_DIR = ROOT / "reports" / "market_demand_os_deep_research_2026-08-25" / "assets"


NAVY = RGBColor(11, 37, 69)
BLUE = RGBColor(46, 116, 181)
DEEP_BLUE = RGBColor(31, 77, 120)
GOLD = RGBColor(191, 139, 38)
INK = RGBColor(32, 38, 45)
GRAY = RGBColor(92, 101, 112)
MUTED = RGBColor(116, 126, 138)
PALE_BLUE = "EAF2F8"
PALE_GOLD = "FFF7E6"
PALE_GRAY = "F2F4F7"
WHITE = RGBColor(255, 255, 255)


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:cs"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D7DBE2", size=6):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_table_layout_fixed(table):
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def set_cell_width(cell, width_inches: float):
    width = Inches(width_inches)
    cell.width = width
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width.twips)))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_field(paragraph):
    paragraph.add_run("Страница ")
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def set_update_fields(document: Document):
    settings = document.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def set_headers_footers(section):
    section.header_distance = Inches(0.42)
    section.footer_distance = Inches(0.42)
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("АЛЮМКОМПЛЕКТ  •  MARKET & DEMAND OS")
    set_run_font(run, size=8.5, color=MUTED, bold=True)
    run = p.add_run("                                      DEEP RESEARCH  •  25.08.2026")
    set_run_font(run, size=8.5, color=MUTED)

    footer = section.footer
    table = footer.add_table(rows=1, cols=2, width=Inches(6.5))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_cell_width(table.cell(0, 0), 4.9)
    set_cell_width(table.cell(0, 1), 1.6)
    left = table.cell(0, 0).paragraphs[0]
    left.paragraph_format.space_after = Pt(0)
    run = left.add_run("Архитектурный release candidate • live-контакт и расходы не разрешены")
    set_run_font(run, size=8, color=MUTED)
    right = table.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right.paragraph_format.space_after = Pt(0)
    add_page_field(right)
    for run in right.runs:
        set_run_font(run, size=8, color=MUTED)
    for cell in table.row_cells(0):
        set_cell_margins(cell, 0, 0, 0, 0)


def configure_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.space_after = Pt(6)
    pf.line_spacing = 1.10
    pf.widow_control = True

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DEEP_BLUE, 8, 4),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.line_spacing = 1.10

    if "Code Block" not in styles:
        style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = styles["Code Block"]
    style.font.name = "Consolas"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    style.font.size = Pt(8.3)
    style.font.color.rgb = INK
    style.paragraph_format.left_indent = Inches(0.16)
    style.paragraph_format.right_indent = Inches(0.16)
    style.paragraph_format.space_before = Pt(4)
    style.paragraph_format.space_after = Pt(7)
    style.paragraph_format.line_spacing = 1.0


def add_hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2E74B5")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([color, underline])
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


INLINE_TOKEN = re.compile(r"(\*\*.+?\*\*|`.+?`|\[[^\]]+\]\([^\)]+\))")


def add_inline_markdown(paragraph, text: str, *, size=None, color=None):
    position = 0
    for match in INLINE_TOKEN.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            set_run_font(run, size=size, color=color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, color=color, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Consolas", size=(size or 10.2), color=DEEP_BLUE)
            r_pr = run._element.get_or_add_rPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), PALE_GRAY)
            r_pr.append(shd)
        else:
            link = re.match(r"\[([^\]]+)\]\(([^\)]+)\)", token)
            if link:
                add_hyperlink(paragraph, link.group(1), link.group(2))
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run, size=size, color=color)


def add_callout(doc: Document, text: str, *, fill=PALE_BLUE, accent="2E74B5"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_layout_fixed(table)
    set_cell_width(table.cell(0, 0), 6.5)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, 160, 190, 160, 190)
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "22")
    left.set(qn("w:color"), accent)
    borders.append(left)
    tc_pr.append(borders)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.12
    add_inline_markdown(p, text, size=11.2, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_metric_strip(doc: Document):
    metrics = [
        ("181", "нормативное\nтребование"),
        ("64", "acceptance-\nсценария"),
        ("11", "machine\nJSON Schema"),
        ("11", "движений\nспроса"),
    ]
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_layout_fixed(table)
    for idx, (value, label) in enumerate(metrics):
        cell = table.cell(0, idx)
        set_cell_width(cell, 1.625)
        set_cell_shading(cell, PALE_GOLD if idx % 2 == 0 else PALE_BLUE)
        set_cell_margins(cell, 100, 90, 100, 90)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        run = p.add_run(value)
        set_run_font(run, size=18, color=NAVY, bold=True)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(0)
        for part_idx, part in enumerate(label.split("\n")):
            if part_idx:
                p2.add_run().add_break()
            run = p2.add_run(part)
            set_run_font(run, size=8.6, color=GRAY, bold=True)
    set_table_borders(table, color="FFFFFF", size=12)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_cover(doc: Document):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(105)
    p.paragraph_format.space_after = Pt(16)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("DEEP RESEARCH  •  RELEASE CANDIDATE")
    set_run_font(run, size=10.5, color=GOLD, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("Market & Demand OS")
    set_run_font(run, size=31, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("Новая архитектура коммерческого спроса АлюмКомплект")
    set_run_font(run, size=15, color=DEEP_BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(30)
    run = p.add_run("Не тендерный бот, а операционная система рынка — от сигнала до оплаты и повтора")
    set_run_font(run, size=10.5, color=GRAY, italic=True)

    line = doc.add_paragraph()
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    line.paragraph_format.space_after = Pt(34)
    run = line.add_run("━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━")
    set_run_font(run, size=8, color=GOLD)

    for label, value in (
        ("Версия", "Исследование 1.0 • AK-MDOS-V7 7.0.0-rc.1"),
        ("Дата", "25 августа 2026"),
        ("Объект", "B2B-производство алюминиевых светопрозрачных конструкций"),
        ("Статус", "Архитектурный release candidate; не разрешение на live-контакт или расходы"),
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(f"{label}: ")
        set_run_font(run, size=9.6, color=GRAY, bold=True)
        run = p.add_run(value)
        set_run_font(run, size=9.6, color=GRAY)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(76)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("Главный принцип: доказанная потребность важнее объёма найденных контактов")
    set_run_font(run, size=10.5, color=NAVY, bold=True)
    doc.add_page_break()


def _load_font(path: str, size: int):
    try:
        return ImageFont.truetype(path, size=size)
    except OSError:
        return ImageFont.load_default()


def _wrap(draw: ImageDraw.ImageDraw, text: str, font, width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        trial = word if not current else f"{current} {word}"
        bbox = draw.textbbox((0, 0), trial, font=font)
        if bbox[2] - bbox[0] <= width:
            current = trial
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def build_architecture_diagram(output: Path):
    output.parent.mkdir(parents=True, exist_ok=True)
    width, height = 1600, 2000
    image = Image.new("RGB", (width, height), "#F8FAFC")
    draw = ImageDraw.Draw(image)
    bold = _load_font(r"C:\Windows\Fonts\arialbd.ttf", 42)
    body = _load_font(r"C:\Windows\Fonts\arial.ttf", 28)
    small = _load_font(r"C:\Windows\Fonts\arial.ttf", 24)
    label = _load_font(r"C:\Windows\Fonts\arialbd.ttf", 26)

    draw.text((110, 60), "MARKET & DEMAND OS v7", font=bold, fill="#0B2545")
    draw.text((110, 120), "Сигналы становятся решениями только через доказательства, policy и outcome", font=body, fill="#5C6570")

    layers = [
        ("01", "SOURCE POLICY PLANE", "лицензия • purpose • consent • retention • quota • provider boundary", "#FFF1D4", "#BF8B26"),
        ("02", "EVENT SIGNAL MESH + EVIDENCE", "raw first • event/observed time • hash • lineage • document intelligence", "#EAF2F8", "#2E74B5"),
        ("03", "BITEMPORAL MARKET WORLD MODEL", "Account • Buying Group • Object/Site • Installed Asset • Interaction • DemandUnit", "#E8F5F1", "#2A7F62"),
        ("04", "DEMAND INTELLIGENCE & ORCHESTRATION", "resolution • fusion • critic • timing • research • motion state • negative evidence", "#EDE9F7", "#6A55A3"),
        ("05", "SELECTIVE DECISION + PORTFOLIO", "legal • Gold profile • capacity • economics • incremental contribution • abstain", "#FDECEF", "#B64B67"),
        ("06", "JOURNEY / ACTION CONTROL", "Bitrix24 • люди • реклама • партнёры • supplier-intent • object protection", "#EAF2F8", "#2E74B5"),
        ("07", "OUTCOME & CAUSAL LEARNING", "ERP • cleared payment • fulfilment • claims • repeat • holdout • calibration • rollback", "#E8F5F1", "#2A7F62"),
    ]
    x0, x1 = 115, 1485
    y = 220
    box_h = 185
    gap = 54
    for idx, (num, title, description, fill, accent) in enumerate(layers):
        draw.rounded_rectangle((x0, y, x1, y + box_h), radius=24, fill=fill, outline=accent, width=4)
        draw.rounded_rectangle((x0 + 20, y + 30, x0 + 105, y + 115), radius=16, fill=accent)
        num_bbox = draw.textbbox((0, 0), num, font=label)
        draw.text((x0 + 62 - (num_bbox[2] - num_bbox[0]) / 2, y + 57), num, font=label, fill="#FFFFFF")
        draw.text((x0 + 135, y + 28), title, font=bold, fill="#0B2545")
        desc_lines = _wrap(draw, description, body, x1 - x0 - 180)
        for line_idx, line in enumerate(desc_lines[:2]):
            draw.text((x0 + 135, y + 92 + line_idx * 38), line, font=body, fill="#39424E")
        if idx < len(layers) - 1:
            arrow_x = width // 2
            draw.line((arrow_x, y + box_h + 5, arrow_x, y + box_h + gap - 12), fill="#7C8794", width=5)
            draw.polygon([(arrow_x - 13, y + box_h + gap - 24), (arrow_x + 13, y + box_h + gap - 24), (arrow_x, y + box_h + gap - 5)], fill="#7C8794")
        y += box_h + gap

    note_y = 1905
    draw.text((115, note_y), "ИИ выдаёт typed claims и proposals — не канонические факты, контакты, цены, оплаты или согласия.", font=small, fill="#B64B67")
    image.save(output, format="PNG", optimize=True)


def add_architecture_figure(doc: Document, image_path: Path):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run()
    inline_shape = run.add_picture(str(image_path), width=Inches(6.35))
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", "Семислойная целевая архитектура Market & Demand OS v7 от Source Policy до causal outcome learning")
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(8)
    run = caption.add_run("Рисунок 1. Целевая архитектура MDOS v7: тендеры и платформы — сенсоры; центр — доказанная единица спроса и outcome loop.")
    set_run_font(run, size=8.8, color=GRAY, italic=True)


def shade_paragraph(paragraph, fill=PALE_GRAY):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def add_code_block(doc: Document, code: str):
    p = doc.add_paragraph(style="Code Block")
    shade_paragraph(p, PALE_GRAY)
    p.paragraph_format.keep_together = True if len(code.splitlines()) <= 12 else False
    for idx, line in enumerate(code.splitlines()):
        if idx:
            p.add_run().add_break()
        run = p.add_run(line)
        set_run_font(run, name="Consolas", size=8.3, color=INK)


def add_table(doc: Document, rows: list[list[str]], heading: str | None = None):
    if len(rows) < 2:
        return
    header = rows[0]
    body = rows[2:] if all(set(cell.strip()) <= {"-", ":"} for cell in rows[1]) else rows[1:]
    landscape = len(header) >= 4 and len(body) >= 8
    if landscape:
        section = doc.add_section(WD_SECTION.NEW_PAGE)
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11)
        section.page_height = Inches(8.5)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.header_distance = Inches(0.35)
        section.footer_distance = Inches(0.35)
        total_width = 9.7
    else:
        total_width = 6.5

    if heading:
        p = doc.add_paragraph(style="Heading 1")
        add_inline_markdown(p, heading)

    table = doc.add_table(rows=1, cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_layout_fixed(table)
    set_table_borders(table)
    widths = None
    if landscape and len(header) == 4:
        widths = [1.15, 2.15, 2.75, 3.65]
    if widths is None:
        widths = [total_width / len(header)] * len(header)
    for idx, cell in enumerate(table.rows[0].cells):
        set_cell_width(cell, widths[idx])
        set_cell_shading(cell, "DCE6F1")
        set_cell_margins(cell, 90, 100, 90, 100)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        add_inline_markdown(p, header[idx], size=8.5 if landscape else 9.2, color=NAVY)
        for run in p.runs:
            run.bold = True
    set_repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])

    for row_idx, row_data in enumerate(body):
        row = table.add_row()
        prevent_row_split(row)
        if row_idx % 2:
            for cell in row.cells:
                set_cell_shading(cell, "FAFBFC")
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell, 78, 96, 78, 96)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            text = row_data[idx] if idx < len(row_data) else ""
            add_inline_markdown(p, text, size=7.7 if landscape else 8.7, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

    if landscape:
        section = doc.add_section(WD_SECTION.NEW_PAGE)
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.top_margin = Inches(0.85)
        section.bottom_margin = Inches(0.78)
        section.header_distance = Inches(0.42)
        section.footer_distance = Inches(0.42)


def starts_block(line: str) -> bool:
    stripped = line.strip()
    return (
        not stripped
        or stripped.startswith("#")
        or stripped.startswith("```")
        or stripped == "---"
        or stripped.startswith("|")
        or re.match(r"^[-*]\s+", stripped) is not None
        or re.match(r"^\d+\.\s+", stripped) is not None
    )


CHAPTER_BREAKS = (
    "1. ",
    "4. ",
    "5. ",
    "6. ",
    "8. ",
    "9. ",
)


def render_markdown(doc: Document, markdown: str, architecture_path: Path):
    lines = markdown.splitlines()
    start = 0
    for idx, line in enumerate(lines):
        if line.strip() == "---":
            start = idx + 1
            break
    lines = lines[start:]
    i = 0
    next_paragraph_callout = False
    architecture_inserted = False
    metric_inserted = False
    pending_table_heading: str | None = None

    while i < len(lines):
        raw = lines[i]
        stripped = raw.strip()
        if not stripped or stripped == "---":
            i += 1
            continue

        if stripped.startswith("## "):
            title = stripped[3:].strip()
            if title.startswith("3. Gap matrix"):
                pending_table_heading = title
                i += 1
                continue
            if any(title.startswith(prefix) for prefix in CHAPTER_BREAKS):
                if doc.paragraphs and doc.paragraphs[-1].text:
                    doc.add_page_break()
            p = doc.add_paragraph(style="Heading 1")
            add_inline_markdown(p, title)
            if title == "Решение в одном абзаце":
                next_paragraph_callout = True
            if title == "Мой окончательный вердикт" and not metric_inserted:
                add_metric_strip(doc)
                metric_inserted = True
            if title.startswith("5. Целевая архитектура") and not architecture_inserted:
                add_architecture_figure(doc, architecture_path)
                architecture_inserted = True
            i += 1
            continue

        if stripped.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            add_inline_markdown(p, stripped[4:].strip())
            i += 1
            continue

        if stripped.startswith("# "):
            i += 1
            continue

        if stripped.startswith("```"):
            code_lines: list[str] = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i].rstrip())
                i += 1
            i += 1
            code = "\n".join(code_lines)
            if "SOURCE POLICY PLANE" in code and architecture_inserted:
                continue
            add_code_block(doc, code)
            continue

        if stripped.startswith("|"):
            table_lines: list[str] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            rows = [[cell.strip() for cell in line.strip("|").split("|")] for line in table_lines]
            add_table(doc, rows, heading=pending_table_heading)
            pending_table_heading = None
            continue

        bullet_match = re.match(r"^[-*]\s+(.+)$", stripped)
        number_match = re.match(r"^(\d+)\.\s+(.+)$", stripped)
        if bullet_match or number_match:
            text = (bullet_match or number_match).group(1 if bullet_match else 2)
            i += 1
            continuation: list[str] = []
            while i < len(lines) and lines[i].strip() and not starts_block(lines[i]):
                continuation.append(lines[i].strip())
                i += 1
            if continuation:
                text += " " + " ".join(continuation)
            if bullet_match:
                p = doc.add_paragraph(style="List Bullet")
                add_inline_markdown(p, text)
            else:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.5)
                p.paragraph_format.first_line_indent = Inches(-0.25)
                p.paragraph_format.space_after = Pt(5)
                p.paragraph_format.line_spacing = 1.10
                number_run = p.add_run(f"{number_match.group(1)}.\t")
                set_run_font(number_run, size=11, color=INK)
                add_inline_markdown(p, text)
            continue

        paragraph_lines = [stripped]
        i += 1
        while i < len(lines) and lines[i].strip() and not starts_block(lines[i]):
            paragraph_lines.append(lines[i].strip())
            i += 1
        text = " ".join(paragraph_lines)
        if next_paragraph_callout:
            add_callout(doc, text)
            next_paragraph_callout = False
        else:
            p = doc.add_paragraph()
            add_inline_markdown(p, text)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.78)
    configure_styles(doc)
    set_headers_footers(section)
    set_update_fields(doc)
    cp = doc.core_properties
    cp.title = "АлюмКомплект Market & Demand OS — глубокое исследование и архитектура"
    cp.subject = "Архитектура коммерческого спроса от сигнала до оплаты и повтора"
    cp.author = "АлюмКомплект / Codex"
    cp.keywords = "lead generation, demand intelligence, Bitrix24, Россия, B2B, алюминиевые конструкции"
    cp.comments = "Architectural release candidate. No live-contact or spend authorization."


def assert_no_placeholders(doc: Document):
    forbidden = ("TODO", "TBD", "PLACEHOLDER", "lorem ipsum")
    text_parts: list[str] = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text_parts.extend(p.text for p in cell.paragraphs)
    joined = "\n".join(text_parts).lower()
    found = [token for token in forbidden if token.lower() in joined]
    if found:
        raise RuntimeError(f"Forbidden placeholders found: {found}")


def build(source: Path, output: Path, asset_dir: Path):
    markdown = source.read_text(encoding="utf-8")
    architecture = asset_dir / "mdos-v7-architecture.png"
    build_architecture_diagram(architecture)
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    render_markdown(doc, markdown, architecture)
    assert_no_placeholders(doc)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(f"DOCX={output}")
    print(f"ARCHITECTURE_ASSET={architecture}")
    print(f"PARAGRAPHS={len(doc.paragraphs)} TABLES={len(doc.tables)} SECTIONS={len(doc.sections)}")


def parse_args():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--asset-dir", type=Path, default=DEFAULT_ASSET_DIR)
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    build(args.source, args.output, args.asset_dir)
