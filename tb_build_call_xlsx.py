# -*- coding: utf-8 -*-
"""Строит Excel «БАЗА ЗВОНИТЬ.xlsx» — остеклители/монтажники из справочников, отсортированные
по приоритету обзвона. Приоритет: тир A (профильные фасад/витраж/алюминий) + ближняя к Ставрополю
зона = первыми (дешёвая логистика → «−30%» сильнее). Пустые колонки под отметки менеджера.
"""
import argparse
import glob
import os
import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from tb_load_prozvon import read_rows, _get, BASE_DIR
from tb_export_calllist import STRONG, INSTALL, ANTI, OKVED, _phones

PROJECT_DIR = Path(__file__).resolve().parent
DEFAULT_OUTPUT_DIR = PROJECT_DIR / "reports"
OUT = DEFAULT_OUTPUT_DIR / "БАЗА ЗВОНИТЬ.xlsx"
OUT2 = DEFAULT_OUTPUT_DIR / "БАЗА ЗВОНИТЬ 1.xlsx"

# Ближняя к Ставрополю зона (лучшая логистика/маржа) — приоритетнее в обзвоне.
SOUTH = ("ставропол", "краснодар", "ростов", "волгоград", "крым", "адыг", "кабардино",
         "карачаево", "осети", "дагестан", "ингуш", "чечен", "калмык", "астрахан", "севастопол")


def _is_south(region, city):
    s = (region + " " + city).lower()
    return any(k in s for k in SOUTH)


def _phone10(s):
    d = re.sub(r"\D", "", str(s or ""))
    return d[-10:] if len(d) >= 10 else ""


def collect(base_dir=BASE_DIR):
    # действующие клиенты из Bitrix — исключаем из обзвона
    try:
        import tb_bitrix_clients
        cli_phones, cli_names = tb_bitrix_clients.client_keys()
        from tb_bitrix_clients import _norm_name
    except Exception as e:
        print(f"  ⚠️ Bitrix недоступен ({e}) — базу собираю БЕЗ вычистки клиентов")
        cli_phones, cli_names = set(), set()

        def _norm_name(_value):
            return ""
    excluded = 0
    seen, rows_out = set(), []
    for f in sorted(glob.glob(os.path.join(os.fspath(base_dir), "export_base_*.xlsx"))):
        try:
            _, rows = read_rows(f)
        except Exception:
            continue
        for r in rows:
            name = _get(r, "название компании", "компания", "название")
            if not name:
                continue
            title = _get(r, "заголовок сайта (title)", "сайт")
            sub = _get(r, "подрубрика")
            okved = _get(r, "главный оквэд (название)")
            okved_code = _get(r, "главный оквэд (код)")
            blob = f"{name} {title} {sub} {okved}"
            if ANTI.search(blob):
                continue
            if STRONG.search(blob):
                tier = "A"
            elif INSTALL.search(blob) or any(okved_code.startswith(c) for c in OKVED):
                tier = "B"
            else:
                continue
            phones = _phones(r)
            if not phones:
                continue
            # вычистка действующих клиентов Bitrix (по телефону или названию)
            ph10 = {_phone10(p) for p in phones.split("  ")}
            if (ph10 & cli_phones) or (_norm_name(name) and _norm_name(name) in cli_names):
                excluded += 1
                continue
            inn = _get(r, "инн")
            key = inn or re.sub(r"\D", "", phones.split("  ")[0]) or name.lower()
            if key in seen:
                continue
            seen.add(key)
            region, city = _get(r, "регион"), _get(r, "город")
            south = _is_south(region, city)
            prio = (1 if tier == "A" and south else 2 if tier == "A"
                    else 3 if south else 4)
            rows_out.append({
                "prio": prio, "Компания": name, "Телефоны": phones, "ИНН": inn,
                "Город": city, "Регион": region,
                "Профиль / сигнал": (title or sub or okved)[:80],
                "ОКВЭД": (okved_code + " " + okved).strip()[:55], "Тир": tier,
            })
    rows_out.sort(key=lambda x: (x["prio"], x["Регион"], x["Компания"]))
    print(f"   🧹 исключено действующих клиентов Bitrix: {excluded}")
    return rows_out


PLABEL = {1: "1 — юг, профильные", 2: "2 — профильные", 3: "3 — юг, монтажники", 4: "4 — монтажники"}
PFILL = {1: "FFC6EFCE", 2: "FFDDF3E0", 3: "FFFFF2CC", 4: "FFF2F2F2"}
COLS = ["Приоритет", "Компания", "Телефоны", "Город", "Регион", "Профиль / сигнал",
        "ОКВЭД", "Тир", "Статус звонка", "Комментарий менеджера", "Дата"]


def _write(data, path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Звонить"
    ws.append(COLS)
    for r in data:
        ws.append([PLABEL[r["prio"]], r["Компания"], r["Телефоны"], r["Город"], r["Регион"],
                   r["Профиль / сигнал"], r["ОКВЭД"], r["Тир"], "", "", ""])
    hdr_fill = PatternFill("solid", fgColor="FF305496")
    for c in range(1, len(COLS) + 1):
        cell = ws.cell(1, c)
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = hdr_fill
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    for i, r in enumerate(data, start=2):
        ws.cell(i, 1).fill = PatternFill("solid", fgColor=PFILL[r["prio"]])
        ws.cell(i, 1).alignment = Alignment(horizontal="center")
        ws.cell(i, 8).alignment = Alignment(horizontal="center")
    for i, w in enumerate([20, 40, 26, 16, 22, 42, 30, 6, 18, 34, 12], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{len(data) + 1}"
    wb.save(path)


def build(out=OUT, out2=OUT2, base_dir=BASE_DIR):
    out = Path(out).expanduser()
    out2 = Path(out2).expanduser()
    if not out.is_absolute():
        out = PROJECT_DIR / out
    if not out2.is_absolute():
        out2 = PROJECT_DIR / out2
    out = out.resolve(strict=False)
    out2 = out2.resolve(strict=False)
    if os.path.normcase(os.fspath(out)) == os.path.normcase(os.fspath(out2)):
        raise ValueError("output paths must be different")
    out.parent.mkdir(parents=True, exist_ok=True)
    out2.parent.mkdir(parents=True, exist_ok=True)
    data = collect(base_dir)
    # делим ПОРОВНУ чередованием (каждому — равная доля горячих приоритетов)
    half_a = data[0::2]   # тебе (файл без цифры)
    half_b = data[1::2]   # партнёру (файл «… 1»)
    _write(half_a, out)
    _write(half_b, out2)
    print(f"✅ {out}  — {len(half_a)} контактов")
    print(f"✅ {out2}  — {len(half_b)} контактов")
    print(f"   Всего разбито: {len(data)} (чередованием — у обоих поровну по приоритетам)")
    for p in (1, 2, 3, 4):
        a = sum(1 for r in half_a if r["prio"] == p)
        b = sum(1 for r in half_b if r["prio"] == p)
        print(f"   Приоритет {p} ({PLABEL[p]}): ты {a} | партнёр {b}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Построить две книги для обзвона")
    parser.add_argument("--dir", default=BASE_DIR, help="папка с export_base_*.xlsx")
    parser.add_argument("--output", default=os.fspath(OUT), help="первая книга")
    parser.add_argument("--output-partner", default=os.fspath(OUT2), help="вторая книга")
    args = parser.parse_args()
    input_dir = Path(args.dir).expanduser()
    if not input_dir.is_absolute():
        input_dir = PROJECT_DIR / input_dir
    input_dir = input_dir.resolve(strict=False)
    if not input_dir.is_dir():
        parser.error(
            f"папка с базами не найдена: {input_dir}. "
            "Передайте --dir или задайте TENDERBOT_PROZVON_DIR."
        )
    build(args.output, args.output_partner, input_dir)
